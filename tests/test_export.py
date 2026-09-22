import csv
import io
import json

from openpyxl import load_workbook

from app.services import export_service as ex
from app.services.formula_guard import is_dangerous, sanitize_cell, sanitize_row

FIELDS = ["name", "email", "amount"]
ROWS = [
    {"name": "Dana", "email": "dana@example.com", "amount": "10"},
    {"name": '=HYPERLINK("http://evil","click")', "email": "e@x.com", "amount": "-5"},
    {"name": "@SUM(A1)", "email": "+cmd|' /C calc'!A0", "amount": "+1.5"},
]
STATUS = ["ready", "error", "warning"]
REPORT = {
    "source_file": "x.csv",
    "schema": "S",
    "rows": 3,
    "ready": 1,
    "warnings": 1,
    "errors": 1,
    "exported_rows": 2,
    "scope": "ready",
    "mapping": [{"source": "=A", "target": "name"}],
    "transformations": [{"kind": "trim", "label": "Trim"}],
    "top_issues": [{"message": "-x", "count": 1}],
}


def test_is_dangerous():
    assert (
        is_dangerous("=1+1")
        and is_dangerous("+cmd")
        and is_dangerous("-cmd")
        and is_dangerous("@SUM")
    )
    assert is_dangerous("\t=x") and is_dangerous("\r=x")
    assert not is_dangerous("-5") and not is_dangerous("+1.5") and not is_dangerous("-0.25e3")
    assert not is_dangerous("+972525551234")  # E.164 phone numbers stay untouched
    assert not is_dangerous("hello") and not is_dangerous("")


def test_sanitize():
    assert sanitize_cell("=SUM(A1)") == "'=SUM(A1)"
    assert sanitize_cell("-5") == "-5"
    assert sanitize_cell(None) is None and sanitize_cell(5) == 5
    assert sanitize_row({"a": "@x", "b": "ok"}) == {"a": "'@x", "b": "ok"}


def test_select_rows_scopes():
    assert [i for i, _ in ex.select_rows(ROWS, STATUS, "ready")] == [0, 2]
    assert [i for i, _ in ex.select_rows(ROWS, STATUS, "all")] == [0, 1, 2]
    assert [i for i, _ in ex.select_rows(ROWS, STATUS, "errors")] == [1]


def test_csv_export_uses_target_names_and_guards_formulas():
    data = ex.to_csv(FIELDS, ROWS)
    assert data.startswith(b"\xef\xbb\xbf")  # BOM for Excel
    text = data.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    assert rows[0] == FIELDS
    assert rows[2][0] == '\'=HYPERLINK("http://evil","click")'
    assert rows[3][0] == "'@SUM(A1)" and rows[3][1] == "'+cmd|' /C calc'!A0"
    assert rows[2][2] == "-5" and rows[3][2] == "+1.5"  # numbers untouched
    for line in text.splitlines()[1:]:
        assert not line.startswith(("=", "+", "-", "@"))


def test_json_export():
    data = json.loads(ex.to_json(FIELDS, ROWS[:1]))
    assert data == [{"name": "Dana", "email": "dana@example.com", "amount": "10"}]
    # JSON is not a spreadsheet: values are kept verbatim there
    assert json.loads(ex.to_json(FIELDS, ROWS[1:2]))[0]["name"].startswith("=")


def test_errors_csv():
    selected = ex.select_rows(ROWS, STATUS, "errors")
    data = ex.errors_csv(FIELDS, selected, {1: [{"field": "email", "message": "Invalid email"}]})
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    assert rows[0] == ["_row", "_issues", *FIELDS]
    assert rows[1][0] == "3" and rows[1][1] == "email: Invalid email"
    assert rows[1][2].startswith("'=")


def test_xlsx_export_two_sheets_and_no_formulas():
    data = ex.to_xlsx(FIELDS, ROWS, REPORT)
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Data", "Import Report"]
    ws = wb["Data"]
    assert [c.value for c in ws[1]] == FIELDS
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            assert cell.data_type != "f", f"formula cell leaked: {cell.value}"
    assert ws["A3"].value == '\'=HYPERLINK("http://evil","click")'
    rep = wb["Import Report"]
    values = [c.value for row in rep.iter_rows() for c in row if c.value is not None]
    assert "Magic Import — Import Report" in values and "'=A" in values and "'-x" in values
    for row in rep.iter_rows():
        for cell in row:
            assert cell.data_type != "f"
