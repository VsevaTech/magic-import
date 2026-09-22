import io

from openpyxl import load_workbook

from tests.conftest import DEMO, make_csv, make_xlsx

CUSTOMER = "Customer Import v1"


def schema_id(client, name=CUSTOMER):
    return next(s["id"] for s in client.get("/api/v1/schemas").json() if s["name"] == name)


def upload(client, filename, data, **form):
    return client.post("/api/v1/imports", files={"file": (filename, data)}, data=form)


def run_to_validated(client, filename, data, mapping_override=None):
    job = upload(client, filename, data).json()
    sid = schema_id(client)
    client.post(f"/api/v1/imports/{job['id']}/schema", json={"schema_id": sid}).raise_for_status()
    r = client.post(
        f"/api/v1/imports/{job['id']}/mapping",
        json={"mapping": mapping_override or {}, "confirm": True},
    )
    r.raise_for_status()
    r = client.post(f"/api/v1/imports/{job['id']}/validate")
    r.raise_for_status()
    return r.json()


# ----------------------------------------------------------------------------- basics
def test_healthz_docs_openapi(client):
    assert client.get("/healthz").json()["status"] == "ok"
    assert client.get("/docs").status_code == 200
    spec = client.get("/openapi.json").json()
    assert spec["info"]["title"] == "Magic Import API"
    paths = spec["paths"]
    for p in [
        "/api/v1/imports",
        "/api/v1/imports/{import_id}",
        "/api/v1/schemas",
        "/api/v1/imports/{import_id}/issues",
        "/api/v1/imports/{import_id}/mapping",
        "/api/v1/imports/{import_id}/validate",
        "/api/v1/imports/{import_id}/export",
    ]:
        assert p in paths, p
    # error contract documented
    assert "ErrorResponse" in spec["components"]["schemas"]
    assert "422" in paths["/api/v1/imports/{import_id}/mapping"]["post"]["responses"]


def test_ui_pages_render(client):
    for path in ["/", "/imports", "/imports/new", "/schemas", "/schemas/new"]:
        r = client.get(path)
        assert r.status_code == 200 and "Magic Import" in r.text, path
    assert client.get("/imports/doesnotexist").status_code == 404
    assert "Page not found" in client.get("/nope").text


def test_error_contract_shape(client):
    r = client.get("/api/v1/imports/nope")
    assert r.status_code == 404
    assert r.json() == {
        "error": {"code": "not_found", "message": "Import 'nope' not found.", "details": []}
    }
    r = client.post("/api/v1/schemas", json={"name": "x"})
    assert r.status_code == 422 and r.json()["error"]["code"] == "validation_error"
    assert r.json()["error"]["details"]
    r = client.get("/api/v1/nothing")
    assert r.status_code == 404 and r.json()["error"]["code"] == "not_found"


def test_builtin_schemas_seeded(client):
    names = {s["name"] for s in client.get("/api/v1/schemas").json()}
    assert {
        "Customer Import v1",
        "Merchant Import",
        "Product Import",
        "Transaction Import",
    } <= names


# ----------------------------------------------------------------------------- upload + security
def test_upload_rejects_bad_extension_and_empty(client):
    r = upload(client, "data.txt", b"a,b\n1,2")
    assert r.status_code == 400 and r.json()["error"]["code"] == "invalid_file"
    r = upload(client, "evil.xlsx.exe", b"MZ")
    assert r.status_code == 400
    r = upload(client, "empty.csv", b"")
    assert r.status_code == 400
    r = upload(client, "bad.xlsx", b"not a workbook")
    assert r.status_code == 400 and r.json()["error"]["code"] == "invalid_file"


def test_upload_rejects_oversized(client, monkeypatch):
    from app import config

    monkeypatch.setenv("MAX_UPLOAD_MB", "1")
    config.get_settings.cache_clear()
    big = b"a,b\n" + b"x,y\n" * 400_000  # ~1.6 MB
    r = upload(client, "big.csv", big)
    assert r.status_code == 413 and r.json()["error"]["code"] == "file_too_large"


def test_path_traversal_filename_is_neutralised(client, tmp_path):
    from app.services.import_service import safe_filename

    for evil in ["../../etc/passwd.csv", "..\\..\\win.csv", "/abs/path/x.csv", "a/../../b.csv"]:
        safe = safe_filename(evil)
        assert "/" not in safe and "\\" not in safe and ".." not in safe
    # multi-sheet xlsx keeps the raw file temporarily -> verify it lands inside the upload dir
    data = make_xlsx(["a"], [["1"]], sheets={"Other": [["b"], ["2"]]})
    r = upload(client, "../../../../evil.xlsx", data)
    assert r.status_code == 201
    from app.config import get_settings

    files = list(get_settings().resolved_upload_dir.iterdir())
    assert len(files) == 1 and files[0].name.endswith("_evil.xlsx")
    assert files[0].resolve().parent == get_settings().resolved_upload_dir.resolve()


def test_malicious_csv_cell_is_stored_verbatim_and_exported_safely(client):
    data = make_csv(
        ["Client", "E-mail Address"],
        [['=HYPERLINK("http://evil","x")', "a@b.com"], ["+cmd", "c@d.com"]],
    )
    job = run_to_validated(client, "evil.csv", data)
    rows = client.get(f"/api/v1/imports/{job['id']}/rows?page_size=10").json()["rows"]
    assert rows[0]["cells"]["customer_name"]["value"].startswith("=HYPERLINK")  # original kept
    csv_out = client.get(f"/api/v1/imports/{job['id']}/export?format=csv&scope=all").content.decode(
        "utf-8-sig"
    )
    assert "'=HYPERLINK" in csv_out and ",'+cmd," in csv_out
    assert "\n=HYPERLINK" not in csv_out and "\n+cmd" not in csv_out
    xlsx = load_workbook(
        io.BytesIO(client.get(f"/api/v1/imports/{job['id']}/export?format=xlsx&scope=all").content)
    )
    assert all(c.data_type != "f" for row in xlsx["Data"].iter_rows() for c in row)


def test_upload_is_deleted_after_processing(client):
    from app.config import get_settings

    r = upload(client, "plain.csv", make_csv(["Client"], [["Dana"]]))
    assert r.status_code == 201
    assert list(get_settings().resolved_upload_dir.iterdir()) == []  # CSV: nothing kept
    data = make_xlsx(["Client"], [["Dana"]], sheets={"Other": [["b"], ["2"]]})
    job = upload(client, "multi.xlsx", data).json()
    assert len(list(get_settings().resolved_upload_dir.iterdir())) == 1  # kept for sheet switching
    client.post(f"/api/v1/imports/{job['id']}/schema", json={"schema_id": schema_id(client)})
    client.post(f"/api/v1/imports/{job['id']}/mapping", json={"mapping": {}, "confirm": True})
    assert (
        list(get_settings().resolved_upload_dir.iterdir()) == []
    )  # deleted once mapping is confirmed


# ----------------------------------------------------------------------------- full flow
def test_full_flow_with_demo_file(client):
    job = upload(client, "acme-customers.xlsx", (DEMO / "acme-customers.xlsx").read_bytes()).json()
    assert job["status"] == "uploaded" and job["row_count"] == 500 and job["column_count"] == 8
    assert job["inspection"]["E-mail Address"]["detected_type"] == "email"
    jid = job["id"]

    m = client.post(f"/api/v1/imports/{jid}/schema", json={"schema_id": schema_id(client)}).json()
    mapping = {x["source_column"]: x for x in m["mapping"]}
    assert (
        mapping["Client"]["target_field"] == "customer_name"
        and mapping["Client"]["confidence"] == "HIGH"
    )
    assert mapping["E-mail Address"]["target_field"] == "email"
    assert mapping["Mobile No."]["target_field"] == "phone"
    assert mapping["Organization"]["target_field"] == "company"
    assert mapping["Country Name"]["target_field"] == "country"
    assert mapping["Joined"]["target_field"] == "created_at"
    assert mapping["Notes"]["confidence"] == "UNMAPPED"

    # duplicate target -> conflict error with machine-readable code
    r = client.post(
        f"/api/v1/imports/{jid}/mapping",
        json={"mapping": {"Organization": "customer_name"}, "confirm": False},
    )
    assert r.status_code == 422 and r.json()["error"]["code"] == "invalid_mapping"
    assert r.json()["error"]["details"][0]["target_field"] == "customer_name"

    # ignore a column, confirm
    r = client.post(
        f"/api/v1/imports/{jid}/mapping", json={"mapping": {"Notes": None}, "confirm": True}
    )
    assert r.status_code == 200
    tr = client.get(f"/api/v1/imports/{jid}/transformations").json()
    kinds = {(t["kind"], t["target_field"]) for t in tr}
    assert (
        ("phone", "phone") in kinds
        and ("country", "country") in kinds
        and ("date", "created_at") in kinds
    )

    job = client.post(f"/api/v1/imports/{jid}/validate").json()
    s = job["summary"]
    assert job["status"] == "validated" and s["rows"] == 500
    assert s["ready"] + s["warning"] + s["error"] == 500
    assert s["error"] >= 40 and s["ready_pct"] > 85
    top = {t["code"] for t in s["top_issues"]}
    assert {
        "email_invalid",
        "duplicate",
        "required_missing",
        "phone_invalid",
        "country_unknown",
    } <= top

    issues = client.get(f"/api/v1/imports/{jid}/issues?severity=error&page_size=5").json()
    assert issues["total"] >= 40 and len(issues["items"]) == 5

    rows = client.get(
        f"/api/v1/imports/{jid}/rows?status=error&code=email_invalid&page_size=1"
    ).json()
    assert rows["total"] >= 5
    row = rows["rows"][0]
    assert row["cells"]["email"]["issues"][0]["code"] == "email_invalid"
    assert (
        row["cells"]["email"]["original"] == row["cells"]["email"]["value"]
    )  # email never "fixed"

    # inline edit revalidates the row and updates the summary
    before_errors = s["error"]
    edited = client.patch(
        f"/api/v1/imports/{jid}/rows/{row['index']}",
        json={"field": "email", "value": "fixed@example.com"},
    ).json()
    assert (
        edited["cells"]["email"]["value"] == "fixed@example.com"
        and edited["cells"]["email"]["edited"]
    )
    assert edited["cells"]["email"]["issues"] == []
    assert edited["cells"]["email"]["original"] != "fixed@example.com"  # original preserved
    after = client.get(f"/api/v1/imports/{jid}").json()["summary"]
    assert after["error"] == before_errors - (1 if edited["status"] != "error" else 0)
    assert after["error"] + after["warning"] + after["ready"] == 500

    # bulk fixes + undo
    fixes = client.get(f"/api/v1/imports/{jid}/bulk-fixes").json()
    assert any(f["type"] == "replace" and f["to_value"] == "IL" for f in fixes)
    fix = next(f for f in fixes if f["to_value"] == "IL")
    r = client.post(
        f"/api/v1/imports/{jid}/bulk-fixes/replace",
        json={"field": "country", "from_value": fix["from_value"], "to_value": "IL"},
    ).json()
    assert r["detail"]["changed"] == fix["count"]
    warn_after_fix = client.get(f"/api/v1/imports/{jid}").json()["summary"]["warning"]
    assert warn_after_fix == after["warning"] - fix["count"]
    undo = client.post(f"/api/v1/imports/{jid}/undo").json()
    assert "IL" in undo["detail"]["undone"]
    assert client.get(f"/api/v1/imports/{jid}").json()["summary"]["warning"] == after["warning"]
    # undo restores an earlier edit of the same cell instead of dropping it
    client.patch(
        f"/api/v1/imports/{jid}/rows/{row['index']}",
        json={"field": "email", "value": "second@example.com"},
    )
    client.post(f"/api/v1/imports/{jid}/undo")
    assert (
        client.get(f"/api/v1/imports/{jid}/rows?page_size=1&search=fixed@example.com").json()[
            "total"
        ]
        == 1
    )
    # reset column drops all manual edits
    r = client.post(f"/api/v1/imports/{jid}/reset-column", json={"field": "email"}).json()
    assert r["detail"]["reset"] == 1
    assert client.get(f"/api/v1/imports/{jid}").json()["summary"]["error"] == before_errors

    # exports
    ready_csv = client.get(f"/api/v1/imports/{jid}/export?format=csv&scope=ready")
    assert ready_csv.status_code == 200 and "text/csv" in ready_csv.headers["content-type"]
    lines = ready_csv.content.decode("utf-8-sig").splitlines()
    assert lines[0] == "customer_id,customer_name,email,phone,company,country,created_at"
    assert len(lines) - 1 == 500 - before_errors
    errors_csv = (
        client.get(f"/api/v1/imports/{jid}/export?format=errors")
        .content.decode("utf-8-sig")
        .splitlines()
    )
    assert errors_csv[0].startswith("_row,_issues,") and len(errors_csv) - 1 == before_errors
    js = client.get(f"/api/v1/imports/{jid}/export?format=json&scope=all").json()
    assert len(js) == 500 and set(js[0]) == {
        "customer_id",
        "customer_name",
        "email",
        "phone",
        "company",
        "country",
        "created_at",
    }
    xlsx = load_workbook(
        io.BytesIO(client.get(f"/api/v1/imports/{jid}/export?format=xlsx&scope=errors").content)
    )
    assert (
        xlsx.sheetnames == ["Data", "Import Report"] and xlsx["Data"].max_row - 1 == before_errors
    )
    report = client.get(f"/api/v1/imports/{jid}/report").json()
    assert report["schema"] == CUSTOMER and report["rows"] == 500 and len(report["mapping"]) == 8
    assert (
        client.get(f"/api/v1/imports/{jid}/report?download=true")
        .headers["content-disposition"]
        .endswith('"import-report.json"')
    )

    done = client.post(f"/api/v1/imports/{jid}/complete").json()
    assert done["status"] == "completed" and done["completed_at"]
    assert client.get("/api/v1/imports").json()["items"][0]["id"] == jid
    assert client.delete(f"/api/v1/imports/{jid}").status_code == 204
    assert client.get(f"/api/v1/imports/{jid}").status_code == 404


def test_three_demo_files_map_to_same_schema(client):
    expectations = {
        "legacy-crm.csv": {
            "Full Name": "customer_name",
            "Mail": "email",
            "Telephone": "phone",
            "Business": "company",
            "Location": "country",
            "Created": "created_at",
        },
        "partner-export.xlsx": {
            "customer": "customer_name",
            "email_address": "email",
            "phone_number": "phone",
            "company_name": "company",
            "country_code": "country",
            "created_at": "created_at",
        },
    }
    for fname, expected in expectations.items():
        job = upload(client, fname, (DEMO / fname).read_bytes()).json()
        m = client.post(
            f"/api/v1/imports/{job['id']}/schema", json={"schema_id": schema_id(client)}
        ).json()
        got = {x["source_column"]: x["target_field"] for x in m["mapping"]}
        assert got == expected, fname
        assert all(x["confidence"] == "HIGH" for x in m["mapping"]), fname


def test_multi_sheet_selection(client):
    data = make_xlsx(
        ["Client", "Mail"],
        [["Dana", "d@x.com"]],
        sheets={"Legacy": [["Full Name", "Telephone"], ["Noa", "0525551234"]]},
    )
    job = upload(client, "multi.xlsx", data).json()
    assert job["sheet_names"] == ["Sheet1", "Legacy"] and job["columns"] == ["Client", "Mail"]
    job = client.post(f"/api/v1/imports/{job['id']}/sheet", json={"sheet_name": "Legacy"}).json()
    assert job["sheet_name"] == "Legacy" and job["columns"] == ["Full Name", "Telephone"]
    r = client.post(f"/api/v1/imports/{job['id']}/sheet", json={"sheet_name": "Nope"})
    assert r.status_code == 400


def test_state_machine_guards(client):
    job = upload(client, "x.csv", make_csv(["Client"], [["Dana"]])).json()
    r = client.post(f"/api/v1/imports/{job['id']}/validate")
    assert r.status_code == 409 and r.json()["error"]["code"] == "invalid_state"
    r = client.get(f"/api/v1/imports/{job['id']}/rows")
    assert r.status_code == 409
    r = client.post(
        f"/api/v1/imports/{job['id']}/mapping", json={"mapping": {"Client": "customer_name"}}
    )
    assert r.status_code == 409  # no schema selected yet


def test_defaults_and_required_unmapped(client):
    data = make_csv(["Client", "Mail"], [["Dana", "d@x.com"], ["Noa", "n@x.com"]])
    job = upload(client, "x.csv", data).json()
    jid = job["id"]
    client.post(f"/api/v1/imports/{jid}/schema", json={"schema_id": schema_id(client)})
    client.post(f"/api/v1/imports/{jid}/mapping", json={"mapping": {}, "confirm": True})
    r = client.put(f"/api/v1/imports/{jid}/defaults", json={"defaults": {"country": "IL"}})
    assert r.status_code == 200
    job = client.post(f"/api/v1/imports/{jid}/validate").json()
    rows = client.get(f"/api/v1/imports/{jid}/rows").json()["rows"]
    assert all(r["cells"]["country"]["value"] == "IL" for r in rows)
    assert job["summary"]["error"] == 0
    r = client.put(f"/api/v1/imports/{jid}/defaults", json={"defaults": {"nope": "x"}})
    assert r.status_code == 422


def test_transformations_api_split_and_combine(client):
    sid = client.post(
        "/api/v1/schemas",
        json={
            "name": "People",
            "fields": [
                {"name": "first_name", "type": "string"},
                {"name": "last_name", "type": "string"},
                {"name": "address", "type": "string"},
                {"name": "full_name", "type": "string", "aliases": ["name"]},
            ],
        },
    ).json()["id"]
    data = make_csv(["Name", "Street", "House", "City"], [["Dana Levi", "Herzl", "12", "Tel Aviv"]])
    job = upload(client, "p.csv", data).json()
    jid = job["id"]
    client.post(f"/api/v1/imports/{jid}/schema", json={"schema_id": sid})
    client.post(
        f"/api/v1/imports/{jid}/mapping", json={"mapping": {"Name": "full_name"}, "confirm": True}
    )
    r = client.put(
        f"/api/v1/imports/{jid}/transformations",
        json={
            "transformations": [
                {"kind": "trim"},
                {
                    "kind": "split",
                    "params": {"source": "full_name", "into": ["first_name", "last_name"]},
                },
                {
                    "kind": "combine",
                    "target_field": "address",
                    "params": {"template": "{Street} {House}, {City}"},
                },
            ]
        },
    )
    assert r.status_code == 200
    client.post(f"/api/v1/imports/{jid}/validate")
    row = client.get(f"/api/v1/imports/{jid}/rows").json()["rows"][0]["cells"]
    assert row["first_name"]["value"] == "Dana" and row["last_name"]["value"] == "Levi"
    assert row["address"]["value"] == "Herzl 12, Tel Aviv"
    r = client.put(
        f"/api/v1/imports/{jid}/transformations", json={"transformations": [{"kind": "bogus"}]}
    )
    assert r.status_code == 422


# ----------------------------------------------------------------------------- schemas
def test_schema_crud_and_validation(client):
    payload = {
        "name": "Terminals",
        "description": "POS terminals",
        "fields": [
            {
                "name": "terminal_id",
                "type": "string",
                "required": True,
                "unique": True,
                "aliases": ["tid"],
            },
            {"name": "model", "type": "enum", "rules": {"enum": ["A920", "A50"]}},
            {"name": "installed_at", "type": "date"},
        ],
    }
    r = client.post("/api/v1/schemas", json=payload)
    assert r.status_code == 201
    s = r.json()
    assert s["fields"][0]["display_name"] == "Terminal id" and s["version"] == 1
    r = client.put(
        f"/api/v1/schemas/{s['id']}",
        json={"fields": payload["fields"] + [{"name": "mcc", "type": "string"}]},
    )
    assert r.status_code == 200 and len(r.json()["fields"]) == 4 and r.json()["version"] == 2
    assert client.get(f"/api/v1/schemas/{s['id']}").json()["name"] == "Terminals"
    # invalid: duplicate names, bad identifier, enum without values, bad type
    bad = [
        {"name": "x", "fields": [{"name": "a"}, {"name": "a"}]},
        {"name": "x", "fields": [{"name": "1bad"}]},
        {"name": "x", "fields": [{"name": "s", "type": "enum"}]},
        {"name": "x", "fields": [{"name": "s", "type": "money"}]},
        {"name": "", "fields": [{"name": "s"}]},
    ]
    for b in bad:
        r = client.post("/api/v1/schemas", json=b)
        assert r.status_code == 422, b
        assert r.json()["error"]["code"] == "validation_error"
    assert client.delete(f"/api/v1/schemas/{s['id']}").status_code == 204
    assert client.get(f"/api/v1/schemas/{s['id']}").status_code == 404


# ----------------------------------------------------------------------------- templates + memory
def test_mapping_templates_and_memory(client):
    sid = schema_id(client)
    # 1st import with odd headers, mapped manually
    data = make_csv(["Cust", "Cust Mail Addr", "Zzz"], [["Dana", "d@x.com", "x"]])
    job = upload(client, "first.csv", data).json()
    m = client.post(f"/api/v1/imports/{job['id']}/schema", json={"schema_id": sid}).json()
    got = {x["source_column"]: x for x in m["mapping"]}
    assert got["Zzz"]["confidence"] == "UNMAPPED"
    client.post(
        f"/api/v1/imports/{job['id']}/mapping",
        json={
            "mapping": {"Cust": "customer_name", "Cust Mail Addr": "email", "Zzz": "company"},
            "confirm": True,
        },
    )
    tpl = client.post(
        f"/api/v1/imports/{job['id']}/mapping/save-template", json={"name": "Odd CRM"}
    ).json()
    assert tpl["mapping"]["Zzz"] == "company"
    assert client.get(f"/api/v1/templates?schema_id={sid}").json()[0]["name"] == "Odd CRM"

    # 2nd import: mapping memory resolves "Zzz" -> company as HIGH/saved, template detected
    job2 = upload(
        client, "second.csv", make_csv(["Cust", "Cust Mail Addr", "Zzz"], [["Noa", "n@x.com", "y"]])
    ).json()
    m2 = client.post(f"/api/v1/imports/{job2['id']}/schema", json={"schema_id": sid}).json()
    got2 = {x["source_column"]: x for x in m2["mapping"]}
    assert (
        got2["Zzz"]["target_field"] == "company"
        and got2["Zzz"]["method"] == "saved"
        and got2["Zzz"]["confidence"] == "HIGH"
    )
    assert (
        m2["template_suggestion"]["name"] == "Odd CRM" and m2["template_suggestion"]["match"] == 1.0
    )
    m3 = client.post(
        f"/api/v1/imports/{job2['id']}/mapping/apply-template", json={"template_id": tpl["id"]}
    ).json()
    assert m3["template_suggestion"] is None
    assert all(x["method"] == "template" for x in m3["mapping"])
    assert client.get("/api/v1/templates").json()[0]["use_count"] == 1
    assert client.delete(f"/api/v1/templates/{tpl['id']}").status_code == 204
    assert client.delete(f"/api/v1/templates/{tpl['id']}").status_code == 404


def test_meta_reports_ai_disabled_without_key(client):
    meta = client.get("/api/v1/meta").json()
    assert meta["ai_available"] is False and "phone" in meta["field_types"]
    assert "email_invalid" in meta["issue_codes"]
