"""Parse CSV / XLSX uploads into headers + rows and profile the columns.

Design rules:
* XLSX is read with openpyxl in read-only, data-only mode: formulas are never evaluated,
  macros/VBA are never touched, only cached cell values are read.
* CSV encoding is detected (UTF-8 / UTF-8-BOM first, then charset-normalizer) with a
  latin-1 fallback so a file never fails on a stray byte.
* Every value is kept as a string (or None for empty) - typing happens later, in the
  transformation engine, where the user can see and undo it.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal

from openpyxl import load_workbook

from app.services.errors import InvalidFileError

ALLOWED_EXTENSIONS = {".csv", ".xlsx"}
DELIMITERS = [",", ";", "\t", "|"]
MAX_SAMPLE_VALUES = 5
_BOM = chr(0xFEFF)


@dataclass
class ParsedFile:
    columns: list[str]
    rows: list[list[str | None]]
    file_type: str
    encoding: str = ""
    delimiter: str | None = None
    sheet_names: list[str] = field(default_factory=list)
    sheet_name: str | None = None

    @property
    def row_count(self) -> int:
        return len(self.rows)

    @property
    def column_count(self) -> int:
        return len(self.columns)


# --------------------------------------------------------------------------- helpers
def _cell_to_str(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, str):
        s = value.strip(_BOM)
        return s if s != "" else None
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, datetime):
        if value.hour == 0 and value.minute == 0 and value.second == 0:
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return repr(value)
    if isinstance(value, int | Decimal):
        return str(value)
    return str(value)


def _dedupe_headers(headers: list[str | None]) -> tuple[list[str], list[str]]:
    """Make headers unique; returns (headers, duplicates_found)."""
    seen: dict[str, int] = {}
    out: list[str] = []
    duplicates: list[str] = []
    for i, raw in enumerate(headers):
        name = (raw or "").strip() or f"Column {i + 1}"
        if name in seen:
            seen[name] += 1
            duplicates.append(name)
            out.append(f"{name} ({seen[name]})")
        else:
            seen[name] = 1
            out.append(name)
    return out, duplicates


def _normalize_row_width(rows: list[list], width: int) -> None:
    for row in rows:
        if len(row) < width:
            row.extend([None] * (width - len(row)))
        elif len(row) > width:
            del row[width:]


def _drop_trailing_empty_rows(rows: list[list]) -> list[list]:
    while rows and all(v is None for v in rows[-1]):
        rows.pop()
    return [r for r in rows if not all(v is None for v in r)]


# --------------------------------------------------------------------------- CSV
def detect_encoding(data: bytes) -> str:
    if data.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"
    try:
        data.decode("utf-8")
        return "utf-8"
    except UnicodeDecodeError:
        pass
    sample = data[:200_000]
    high = sum(1 for b in sample if b >= 0x80)
    if high / max(1, len(sample)) < 0.1:
        # Mostly ASCII with a few accented characters: almost always a Windows-1252 export.
        try:
            sample.decode("cp1252")
            return "cp1252"
        except UnicodeDecodeError:
            pass
    try:
        from charset_normalizer import from_bytes

        best = from_bytes(sample).best()
        if best and best.encoding:
            return best.encoding
    except Exception:  # pragma: no cover - defensive
        pass
    return "latin-1"


def detect_delimiter(sample: str) -> str:
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters="".join(DELIMITERS))
        return dialect.delimiter
    except csv.Error:
        pass
    first_lines = sample.splitlines()[:20]
    counts = {d: sum(line.count(d) for line in first_lines) for d in DELIMITERS}
    best = max(counts, key=counts.get)
    return best if counts[best] > 0 else ","


def parse_csv(data: bytes) -> ParsedFile:
    if not data.strip():
        raise InvalidFileError("The file is empty.")
    encoding = detect_encoding(data)
    try:
        text = data.decode(encoding, errors="strict")
    except (UnicodeDecodeError, LookupError):
        encoding = "latin-1"
        text = data.decode("latin-1", errors="replace")
    text = text.lstrip(_BOM)
    if "\x00" in text[:4096]:
        raise InvalidFileError("The file does not look like a text CSV file.")
    delimiter = detect_delimiter(text[:20_000])
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    raw_rows = list(reader)
    raw_rows = [r for r in raw_rows if any((c or "").strip() for c in r)]
    if not raw_rows:
        raise InvalidFileError("The file contains no data rows.")
    headers, _dups = _dedupe_headers([h.strip() for h in raw_rows[0]])
    rows = [[_cell_to_str(c) for c in r] for r in raw_rows[1:]]
    _normalize_row_width(rows, len(headers))
    rows = _drop_trailing_empty_rows(rows)
    return ParsedFile(
        columns=headers,
        rows=rows,
        file_type="csv",
        encoding=encoding.upper().replace("_", "-"),
        delimiter=delimiter,
    )


# --------------------------------------------------------------------------- XLSX
def list_sheets(data: bytes) -> list[str]:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_vba=False)
    except Exception as exc:  # zipfile / openpyxl errors
        raise InvalidFileError("The file is not a valid XLSX workbook.") from exc
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def parse_xlsx(data: bytes, sheet_name: str | None = None) -> ParsedFile:
    if not data:
        raise InvalidFileError("The file is empty.")
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_vba=False)
    except Exception as exc:
        raise InvalidFileError("The file is not a valid XLSX workbook.") from exc
    try:
        sheet_names = list(wb.sheetnames)
        if not sheet_names:
            raise InvalidFileError("The workbook has no sheets.")
        if sheet_name is None:
            sheet_name = sheet_names[0]
        if sheet_name not in sheet_names:
            raise InvalidFileError(f"Sheet '{sheet_name}' does not exist in the workbook.")
        ws = wb[sheet_name]
        raw_rows = [[_cell_to_str(c) for c in row] for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    raw_rows = [r for r in raw_rows if any(v is not None for v in r)]
    if not raw_rows:
        raise InvalidFileError(f"Sheet '{sheet_name}' contains no data.")
    headers, _dups = _dedupe_headers(raw_rows[0])
    # drop fully empty trailing header columns
    while (
        headers
        and headers[-1].startswith("Column ")
        and all((len(r) < len(headers) or r[len(headers) - 1] is None) for r in raw_rows[1:])
    ):
        headers.pop()
    rows = raw_rows[1:]
    _normalize_row_width(rows, len(headers))
    rows = _drop_trailing_empty_rows(rows)
    return ParsedFile(
        columns=headers,
        rows=rows,
        file_type="xlsx",
        encoding="UTF-8",
        sheet_names=sheet_names,
        sheet_name=sheet_name,
    )


def parse_upload(filename: str, data: bytes, sheet_name: str | None = None) -> ParsedFile:
    ext = extension_of(filename)
    if ext == ".csv":
        return parse_csv(data)
    if ext == ".xlsx":
        return parse_xlsx(data, sheet_name)
    raise InvalidFileError("Only CSV and XLSX files are supported.")


def extension_of(filename: str) -> str:
    name = filename.lower().rsplit("/", 1)[-1].rsplit("\\", 1)[-1]
    if "." not in name:
        return ""
    return "." + name.rsplit(".", 1)[-1]


# --------------------------------------------------------------------------- inspection
_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
_PHONE_RE = re.compile(r"^\+?[\d\s\-().]{7,20}$")
_INT_RE = re.compile(r"^[+-]?\d{1,18}$")
_DEC_RE = re.compile(r"^[+-]?[\d.,\s]*\d[\d.,]*$")
_DATE_RE = re.compile(
    r"^(\d{4}-\d{2}-\d{2}|\d{1,2}[./-]\d{1,2}[./-]\d{2,4}|"
    r"[A-Za-z]{3,9}\.? \d{1,2},? \d{4}|\d{1,2} [A-Za-z]{3,9}\.? \d{4})"
    r"([ T]\d{1,2}:\d{2}(:\d{2})?)?$"
)
_URL_RE = re.compile(r"^https?://", re.I)
_BOOL_VALUES = {"true", "false", "yes", "no", "y", "n", "1", "0"}
_COUNTRY_CODE_RE = re.compile(r"^[A-Za-z]{2,3}$")


def guess_type(values: list[str]) -> str:
    """Guess a probable data type from non-empty sample values."""
    if not values:
        return "empty"
    n = len(values)

    def ratio(pred) -> float:
        return sum(1 for v in values if pred(v)) / n

    if ratio(lambda v: bool(_EMAIL_RE.match(v))) >= 0.8:
        return "email"
    if ratio(lambda v: v.lower() in _BOOL_VALUES) >= 0.95:
        return "boolean"
    if ratio(lambda v: bool(_URL_RE.match(v))) >= 0.8:
        return "url"
    if ratio(lambda v: bool(_DATE_RE.match(v))) >= 0.8:
        return "date"
    if ratio(lambda v: bool(_INT_RE.match(v))) >= 0.9:
        return "integer"
    if ratio(lambda v: bool(_PHONE_RE.match(v)) and sum(c.isdigit() for c in v) >= 7) >= 0.8:
        return "phone"
    if ratio(lambda v: bool(_DEC_RE.match(v))) >= 0.9:
        return "decimal"
    if ratio(lambda v: bool(_COUNTRY_CODE_RE.match(v))) >= 0.9:
        return "country_code"
    return "string"


def inspect(parsed: ParsedFile) -> dict:
    """Profile every column: nulls, uniqueness, examples and probable type."""
    profile: dict[str, dict] = {}
    total = parsed.row_count
    duplicates = _duplicate_headers(parsed.columns)
    for idx, col in enumerate(parsed.columns):
        values = [r[idx] for r in parsed.rows]
        non_empty = [v for v in values if v is not None and v.strip() != ""]
        distinct = set(non_empty)
        examples: list[str] = []
        for v in non_empty:
            if v not in examples:
                examples.append(v)
            if len(examples) >= MAX_SAMPLE_VALUES:
                break
        sample_for_type = non_empty[:500]
        profile[col] = {
            "index": idx,
            "empty_pct": round(100 * (total - len(non_empty)) / total, 1) if total else 100.0,
            "unique_pct": round(100 * len(distinct) / len(non_empty), 1) if non_empty else 0.0,
            "distinct_count": len(distinct),
            "non_empty_count": len(non_empty),
            "is_empty": len(non_empty) == 0,
            "is_duplicate_header": col in duplicates,
            "detected_type": guess_type(sample_for_type),
            "examples": examples,
        }
    return profile


def _duplicate_headers(columns: list[str]) -> set[str]:
    base = [re.sub(r" \(\d+\)$", "", c) for c in columns]
    return {c for c in columns if base.count(re.sub(r" \(\d+\)$", "", c)) > 1}
