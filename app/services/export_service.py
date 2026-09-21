"""Export normalized data as CSV / XLSX / JSON plus an import report."""

from __future__ import annotations

import csv
import io
import json
from datetime import UTC, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.formula_guard import sanitize_cell

SCOPES = ("ready", "all", "errors")


def select_rows(rows: list[dict], statuses: list[str], scope: str) -> list[tuple[int, dict]]:
    if scope not in SCOPES:
        raise ValueError(f"Unknown export scope: {scope}")
    out: list[tuple[int, dict]] = []
    for i, rec in enumerate(rows):
        status = statuses[i] if i < len(statuses) else "ready"
        if (
            scope == "all"
            or scope == "ready"
            and status != "error"
            or scope == "errors"
            and status == "error"
        ):
            out.append((i, rec))
    return out


def to_csv(field_names: list[str], rows: list[dict]) -> bytes:
    buf = io.StringIO(newline="")
    writer = csv.writer(buf, lineterminator="\r\n")
    writer.writerow([sanitize_cell(n) for n in field_names])
    for rec in rows:
        writer.writerow(
            [sanitize_cell(rec.get(n)) if rec.get(n) is not None else "" for n in field_names]
        )
    return buf.getvalue().encode("utf-8-sig")


def to_json(field_names: list[str], rows: list[dict]) -> bytes:
    data = [{n: rec.get(n) for n in field_names} for rec in rows]
    return json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8")


def errors_csv(
    field_names: list[str], rows: list[tuple[int, dict]], issues_by_row: dict[int, list]
) -> bytes:
    """Error rows plus two helper columns: row number and the list of problems."""
    buf = io.StringIO(newline="")
    writer = csv.writer(buf, lineterminator="\r\n")
    writer.writerow(["_row", "_issues", *field_names])
    for i, rec in rows:
        problems = "; ".join(
            f"{iss['field']}: {iss['message']}" if iss.get("field") else iss["message"]
            for iss in issues_by_row.get(i, [])
        )
        writer.writerow(
            [
                i + 2,
                sanitize_cell(problems),
                *[sanitize_cell(rec.get(n)) if rec.get(n) is not None else "" for n in field_names],
            ]
        )
    return buf.getvalue().encode("utf-8-sig")


_HEADER_FILL = PatternFill("solid", fgColor="1E293B")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def to_xlsx(field_names: list[str], rows: list[dict], report: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(field_names)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for rec in rows:
        ws.append([None] * len(field_names))
        r = ws.max_row
        for c, name in enumerate(field_names, start=1):
            value = sanitize_cell(rec.get(name))
            cell = ws.cell(row=r, column=c)
            cell.value = value
            if isinstance(value, str):
                cell.data_type = "s"  # never let openpyxl interpret the value as a formula
    for c, name in enumerate(field_names, start=1):
        width = max(
            12,
            min(
                40,
                max((len(str(rec.get(name) or "")) for rec in rows[:200]), default=0) + 2,
                len(name) + 2,
            ),
        )
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = "A2"

    rep = wb.create_sheet("Import Report")
    rep.column_dimensions["A"].width = 28
    rep.column_dimensions["B"].width = 60
    rep.append(["Magic Import — Import Report", ""])
    rep["A1"].font = Font(bold=True, size=14)
    rep.append(["Generated at", report.get("generated_at", "")])
    rep.append(["Source file", report.get("source_file", "")])
    rep.append(["Schema", report.get("schema", "")])
    rep.append(["Rows in file", report.get("rows", 0)])
    rep.append(["Ready", report.get("ready", 0)])
    rep.append(["Warnings", report.get("warnings", 0)])
    rep.append(["Errors", report.get("errors", 0)])
    rep.append(["Exported rows", report.get("exported_rows", 0)])
    rep.append(["Export scope", report.get("scope", "")])
    rep.append([])
    rep.append(["Mapping", ""])
    rep.cell(row=rep.max_row, column=1).font = Font(bold=True)
    for m in report.get("mapping", []):
        rep.append([sanitize_cell(m["source"]), sanitize_cell(m["target"] or "— ignored —")])
    rep.append([])
    rep.append(["Transformations", ""])
    rep.cell(row=rep.max_row, column=1).font = Font(bold=True)
    for t in report.get("transformations", []):
        rep.append([sanitize_cell(t.get("kind", "")), sanitize_cell(t.get("label", ""))])
    rep.append([])
    rep.append(["Top issues", ""])
    rep.cell(row=rep.max_row, column=1).font = Font(bold=True)
    for iss in report.get("top_issues", []):
        rep.append([sanitize_cell(iss["message"]), iss["count"]])
    for row in rep.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build_report(
    job,
    summary: dict,
    scope: str,
    exported_rows: int,
    mapping: list[dict],
    transformations: list[dict],
    top_issues: list[dict],
) -> dict:
    return {
        "generated_at": datetime.now(UTC).replace(microsecond=0).isoformat(),
        "import_id": job.id,
        "source_file": job.original_filename,
        "schema": job.schema.name if job.schema else None,
        "rows": job.row_count,
        "ready": summary.get("ready", 0),
        "warnings": summary.get("warning", 0),
        "errors": summary.get("error", 0),
        "ready_pct": summary.get("ready_pct", 0),
        "scope": scope,
        "exported_rows": exported_rows,
        "mapping": mapping,
        "transformations": transformations,
        "top_issues": top_issues,
    }
